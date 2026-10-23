#!/usr/bin/env python3
"""Analyze a synthetic 500k-item inventory population and build PT530 files."""
from __future__ import annotations
import argparse
import csv
import gzip
import heapq
import json
import math
import random
import statistics
from datetime import date
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

N_ASSETS = 13_219_635_000
INVENTORY_TOTAL = 6_951_508_000
LAND_TOTAL = 5_678_238_000
LAND_NONCURRENT = 4_864_673_000
LAND_CURRENT = LAND_TOTAL - LAND_NONCURRENT
WIP_TOTAL = 1_174_207_000
COMPLETED_TOTAL = 99_063_000
INVENTORY_CURRENT = 2_086_835_000
INVENTORY_NONCURRENT = 4_864_673_000
REVENUE = 4_343_008_000
PBT = 1_078_495_000
ASSETS_2024 = 10_067_406_000
INVENTORY_2024 = 5_532_090_000
LAND_2024 = 4_587_823_000
REVENUE_2024 = 3_348_470_000
PBT_2024 = 792_298_000
POPULATION_CUTOFF = date(2025, 12, 31)
CLASSIFICATION_HORIZON = date(2026, 12, 31)
STALE_LIMIT = date(2023, 1, 1)
ACCESS_DATE = "2026-10-22"
SOURCE_CVM = "DFP Direcional Engenharia S.A., 31/12/2025, V1, CVM protocolo 021350DFP311220250100155189-71, consulta documental NSD 155189, acesso 22/10/2026; BP pp.16–17, Nota 5 p.69."
SOURCE_RI = "RI Direcional, DFP 4T25 V1, PDF direto: https://api.mziq.com/mzfilemanager/v2/d/ada9bc2c-f7d0-4359-9eaf-851b679ab788/184f5357-c3a8-12a9-3a85-cdb268303bcb?origin=2"
CLASS_BALANCES = {
    "TERRENOS_A_INCORPORAR": LAND_TOTAL,
    "IMOVEIS_EM_CONSTRUCAO": WIP_TOTAL,
    "IMOVEIS_CONCLUIDOS": COMPLETED_TOTAL,
}
FIELDS = [
    "synthetic_item_id", "synthetic_contract_id", "synthetic_project_id",
    "synthetic_registry_reference", "inventory_class", "item_type", "uf_code",
    "municipality_code_synthetic", "area_m2", "acquisition_date",
    "planned_realization_date", "classification_at_2025_12_31", "title_status",
    "title_registration_date", "valuation_date", "historical_cost_cents_synthetic",
    "book_value_cents_synthetic", "estimated_nrv_cents_synthetic",
    "injected_exception_code", "synthetic_record",
]
RULES = ["TITLE_DOCUMENT_INCOMPLETE", "AREA_INVALID", "NRV_BELOW_BOOK", "CLASSIFICATION_MISMATCH", "VALUATION_STALE"]


def parse_date(s: str) -> date:
    return date.fromisoformat(s)


def format_brl(cents: int | float) -> str:
    return f"R$ {cents / 100:,.2f}"


def detected_rules(row: dict[str, str]) -> list[str]:
    flags = []
    book = int(row["book_value_cents_synthetic"])
    nrv = int(row["estimated_nrv_cents_synthetic"])
    if row["title_status"] != "DOCUMENTADO" or not row["title_registration_date"]:
        flags.append("TITLE_DOCUMENT_INCOMPLETE")
    if int(row["area_m2"]) <= 0:
        flags.append("AREA_INVALID")
    if nrv < book:
        flags.append("NRV_BELOW_BOOK")
    planned = parse_date(row["planned_realization_date"])
    current = row["classification_at_2025_12_31"] == "CIRCULANTE"
    if (planned <= CLASSIFICATION_HORIZON and not current) or (planned > CLASSIFICATION_HORIZON and current):
        flags.append("CLASSIFICATION_MISMATCH")
    if parse_date(row["valuation_date"]) < STALE_LIMIT:
        flags.append("VALUATION_STALE")
    return flags


def style_sheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="17365D")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 32
    for col in ws.columns:
        vals = list(col)
        width = min(48, max(12, max((len(str(c.value or "")) for c in vals[:200]), default=10) + 2))
        ws.column_dimensions[get_column_letter(vals[0].column)].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(vertical="top", wrap_text=True)


def analyze(dataset: Path, outdir: Path, sample_size: int = 1000, seed: int = 20251232):
    if not 1 <= sample_size <= 500_000:
        raise ValueError("sample-size must be between 1 and 500000")
    outdir.mkdir(parents=True, exist_ok=True)
    sample_rng = random.Random(seed)
    sample: list[dict[str, str]] = []
    top_heap: list[tuple[int, str, dict[str, str]]] = []
    n = total_cents = total_cost_cents = current_cents = noncurrent_cents = 0
    class_cents = {k: 0 for k in CLASS_BALANCES}
    class_counts = {k: 0 for k in CLASS_BALANCES}
    injected_ids: dict[str, set[str]] = {k: set() for k in RULES}
    detected_ids: dict[str, set[str]] = {k: set() for k in RULES}
    detected_rows: list[dict[str, str]] = []
    sample_nrv_shortfalls_brl: list[float] = []
    seen_ids: set[str] = set()
    opener = gzip.open if dataset.suffix == ".gz" else open
    with opener(dataset, "rt", newline="", encoding="utf-8") as f:
        for row in csv.DictReader(f):
            n += 1
            item_id = row["synthetic_item_id"]
            if item_id in seen_ids: raise ValueError(f"Duplicate synthetic item ID: {item_id}")
            seen_ids.add(item_id)
            book = int(row["book_value_cents_synthetic"])
            cost = int(row["historical_cost_cents_synthetic"])
            nrv = int(row["estimated_nrv_cents_synthetic"])
            total_cents += book; total_cost_cents += cost
            cls = row["inventory_class"]
            if cls not in class_cents: raise ValueError(f"Unknown inventory class: {cls}")
            class_cents[cls] += book; class_counts[cls] += 1
            if row["classification_at_2025_12_31"] == "CIRCULANTE": current_cents += book
            else: noncurrent_cents += book
            injected = row["injected_exception_code"]
            if injected in injected_ids: injected_ids[injected].add(item_id)
            flags = detected_rules(row)
            for rule in flags: detected_ids[rule].add(item_id)
            if flags:
                detected = dict(row); detected["detected_exception_codes"] = ";".join(flags)
                detected_rows.append(detected)
            if len(sample) < sample_size: sample.append(row.copy())
            else:
                j = sample_rng.randrange(n)
                if j < sample_size: sample[j] = row.copy()
            key = (book, item_id, row.copy())
            if len(top_heap) < 200: heapq.heappush(top_heap, key)
            elif book > top_heap[0][0]: heapq.heapreplace(top_heap, key)

    if n != 500_000: raise ValueError(f"Expected 500,000 rows; observed {n:,}.")
    if len(seen_ids) != n: raise ValueError("Synthetic item IDs are not unique.")
    if total_cents != INVENTORY_TOTAL * 100:
        raise ValueError(f"Reconciliation failed: observed cents {total_cents}; target {INVENTORY_TOTAL * 100}.")
    if total_cost_cents != total_cents: raise ValueError("Synthetic cost and carrying values do not reconcile.")
    for cls, amount in CLASS_BALANCES.items():
        if class_cents[cls] != amount * 100:
            raise ValueError(f"Class reconciliation failed for {cls}: {class_cents[cls]} != {amount * 100}.")
    if current_cents != INVENTORY_CURRENT * 100 or noncurrent_cents != INVENTORY_NONCURRENT * 100:
        raise ValueError("Current/noncurrent synthetic totals do not reconcile to public balances.")

    for row in sample:
        shortfall = max(0, int(row["book_value_cents_synthetic"]) - int(row["estimated_nrv_cents_synthetic"])) / 100
        sample_nrv_shortfalls_brl.append(shortfall)
        row["detected_exception_codes"] = ";".join(detected_rules(row))
        row["sample_method"] = f"SRS sem reposição; seed {seed}"
    sample_shortfall_sum = sum(sample_nrv_shortfalls_brl)
    projected_shortfall = sample_shortfall_sum * n / len(sample)
    z_95, expected_attribute_rate, tolerable_precision = 1.96, 0.50, 0.05
    n0_attribute = z_95 ** 2 * expected_attribute_rate * (1 - expected_attribute_rate) / tolerable_precision ** 2
    n_required_attribute = n0_attribute / (1 + (n0_attribute - 1) / n)
    n_required_attribute_ceiling = math.ceil(n_required_attribute)
    if len(sample_nrv_shortfalls_brl) > 1:
        sd = statistics.stdev(sample_nrv_shortfalls_brl)
        se = sd / math.sqrt(len(sample)) * math.sqrt((n - len(sample)) / (n - 1))
        margin = 1.96 * se * n
    else: margin = 0.0

    top_rows = [x[2] for x in sorted(top_heap, reverse=True)]
    targeted_ids = {r["synthetic_item_id"] for r in top_rows}
    random_ids = {r["synthetic_item_id"] for r in sample}
    exception_summary = {}
    for rule in RULES:
        planted, detected = injected_ids[rule], detected_ids[rule]
        exception_summary[rule] = {
            "injected": len(planted), "detected": len(detected), "missed": len(planted - detected),
            "unexpected": len(detected - planted), "exact_match": planted == detected,
            "example_ids": sorted(detected)[:10],
        }
    materiality = {
        "overall_benchmark": "Ativo total consolidado em 31/12/2025", "overall_percent": 0.0075,
        "overall_brl": round(N_ASSETS * 0.0075), "performance_percent_of_overall": 0.70,
        "performance_brl": round(N_ASSETS * 0.0075 * 0.70),
        "specific_benchmark": "Estoques consolidados, Nota 5, 2025", "specific_percent": 0.005,
        "specific_brl": round(INVENTORY_TOTAL * 0.005), "clearly_trivial_percent": 0.05,
        "clearly_trivial_brl": round(N_ASSETS * 0.0075 * 0.05),
        "note": "Percentuais são premissas didáticas de julgamento, não parâmetros normativos nem materialidade estabelecida pelo auditor independente da companhia.",
    }
    summary = {
        "synthetic_only": True, "source_date": ACCESS_DATE, "population_rows": n,
        "unique_synthetic_item_ids": len(seen_ids), "sum_book_value_brl": total_cents / 100,
        "public_balance_inventory_brl": INVENTORY_TOTAL,
        "reconciliation_difference_brl": total_cents / 100 - INVENTORY_TOTAL,
        "public_balance_class_brl": CLASS_BALANCES,
        "synthetic_class_brl": {k: v / 100 for k, v in class_cents.items()},
        "class_record_counts": class_counts,
        "public_balance_current_brl": INVENTORY_CURRENT, "public_balance_noncurrent_brl": INVENTORY_NONCURRENT,
        "current_amount_synthetic_brl": current_cents / 100, "noncurrent_amount_synthetic_brl": noncurrent_cents / 100,
        "land_total_brl": LAND_TOTAL, "land_current_derived_brl": LAND_CURRENT,
        "land_noncurrent_brl": LAND_NONCURRENT, "land_share_of_inventory_pct": LAND_TOTAL / INVENTORY_TOTAL * 100,
        "land_share_of_total_assets_pct": LAND_TOTAL / N_ASSETS * 100, "materiality": materiality,
        "sample_method": "Amostragem aleatória simples sem reposição de itens, algoritmo de reservatório, seed fixo; base integralmente sintética.",
        "sample_size": len(sample), "sample_population_fraction_pct": len(sample) / n * 100, "sample_seed": seed,
        "sample_size_basis": {
            "design": "SRS without replacement; illustrative proportion estimate for binary attributes",
            "confidence": 0.95, "z": z_95, "expected_proportion_conservative": expected_attribute_rate,
            "tolerable_precision_percentage_points": tolerable_precision * 100,
            "initial_n0": n0_attribute, "finite_population_corrected_n": n_required_attribute,
            "minimum_rounded_up": n_required_attribute_ceiling, "actual_n": len(sample),
            "coverage_multiple": len(sample) / n_required_attribute_ceiling,
            "limitation": "Tamanho ilustrativo em base sintética; não substitui desenho real, não assegura detecção de distorções monetárias e não permite extrapolação para a companhia.",
        },
        "sample_detected_exception_counts": {r: sum(1 for row in sample if r in detected_rules(row)) for r in RULES},
        "targeted_top_value_size": len(top_rows), "targeted_top_value_unique_count": len(targeted_ids),
        "overlap_srs_and_targeted": len(targeted_ids & random_ids),
        "projected_synthetic_nrv_shortfall_brl": projected_shortfall,
        "projected_synthetic_nrv_shortfall_ci95_brl": [max(0, projected_shortfall - margin), projected_shortfall + margin],
        "actual_injected_synthetic_nrv_shortfall_brl": sum(max(0, int(r["book_value_cents_synthetic"]) - int(r["estimated_nrv_cents_synthetic"])) for r in detected_rows if r["injected_exception_code"] == "NRV_BELOW_BOOK") / 100,
        "population_exception_results": exception_summary,
        "sources": {"cvm_dfp": SOURCE_CVM, "ri_dfp": SOURCE_RI},
        "scope_limit": "Nenhuma exceção da base é uma constatação da Direcional; todos os itens e exceções são sintéticos. A projeção não se aplica aos saldos reais.",
    }
    (outdir / "PT530_resultados.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    def write_csv(path: Path, rows: list[dict[str, str]], extras: list[str] | None = None):
        if rows:
            fields = FIELDS + (extras or [])
            with path.open("w", newline="", encoding="utf-8") as f:
                w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
                w.writeheader(); w.writerows(rows)
    write_csv(outdir / "amostra_aleatoria_1000.csv", sample, ["detected_exception_codes", "sample_method"])
    write_csv(outdir / "amostra_direcionada_top200.csv", top_rows, ["detected_exception_codes"])
    write_csv(outdir / "excecoes_detectadas_sinteticas.csv", detected_rows, ["detected_exception_codes"])

    dictionary = [
        ["synthetic_item_id", "string", "Chave primária única de item/posição de custo inteiramente sintética.", "SYN-INV-0000001…"],
        ["synthetic_contract_id", "string", "Contrato aquisitivo fictício, exclusivo por linha para evitar repetição de registros.", "SYN-CON-0000001…"],
        ["synthetic_project_id", "string", "Projeto fictício; repetição entre itens do mesmo projeto é legítima.", "SYN-PRJ-0001…"],
        ["synthetic_registry_reference", "string", "Referência registral fictícia; não corresponde a matrícula real.", "SYN-REG-0000001…"],
        ["inventory_class", "enum", "Classe conforme Nota 5; valores individuais e IDs são artificiais.", "TERRENOS_A_INCORPORAR / IMOVEIS_EM_CONSTRUCAO / IMOVEIS_CONCLUIDOS"],
        ["item_type", "enum", "Natureza sintética do item individual da rubrica Estoques.", "PARCELA_DE_TERRENO / UNIDADE_EM_CONSTRUCAO / UNIDADE_CONCLUIDA"],
        ["uf_code", "string", "UF atribuída artificialmente; não identifica localização real.", "MG/SP/RJ/…"],
        ["municipality_code_synthetic", "string", "Código municipal sintético, sem correspondência com município real.", "SYN-MUN-0001…"],
        ["area_m2", "integer", "Área artificial em m²; anomalias de área zero foram implantadas.", "0 ou inteiro positivo"],
        ["acquisition_date", "ISO date", "Data de aquisição inventada entre 2018 e 2024.", "AAAA-MM-DD"],
        ["planned_realization_date", "ISO date", "Data fictícia para testar expectativa de realização e classificação.", "AAAA-MM-DD"],
        ["classification_at_2025_12_31", "enum", "Classificação corrente/não corrente sintética, sujeita a inconsistência plantada.", "CIRCULANTE/NAO_CIRCULANTE"],
        ["title_status", "enum", "Evidência dominial fictícia para itens imobiliários.", "DOCUMENTADO/PENDENTE"],
        ["title_registration_date", "ISO date/text", "Data fictícia; vazia em casos plantados como pendentes.", "AAAA-MM-DD ou vazio"],
        ["valuation_date", "ISO date", "Data fictícia para teste de atualidade da mensuração.", "AAAA-MM-DD"],
        ["historical_cost_cents_synthetic", "integer cents", "Custo alocado artificialmente; soma reconcilia por classe e total.", "centavos de BRL"],
        ["book_value_cents_synthetic", "integer cents", "Valor contábil sintético; soma exatamente R$6.951.508.000,00.", "centavos de BRL"],
        ["estimated_nrv_cents_synthetic", "integer cents", "Estimativa sintética de VLR; há exceções de teste implantadas.", "centavos de BRL"],
        ["injected_exception_code", "enum/text", "Exceção implantada pelo gerador; vazio se não plantada.", "; ou vazio"],
        ["synthetic_record", "string", "Rótulo explícito de artificialidade em todas as linhas.", "YES — DADOS SINTÉTICOS"],
    ]
    with (outdir / "dicionario_variaveis.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f); w.writerow(["variavel", "tipo", "descricao", "dominio/unidade"]); w.writerows(dictionary)

    wb = Workbook(); wb.calculation.fullCalcOnLoad = True; wb.calculation.forceFullCalc = True; wb.calculation.calcMode = "auto"
    ws = wb.active; ws.title = "Materialidade"
    ws.append(["Parâmetro", "Base R$", "Percentual", "Valor R$ (fórmula Excel)", "Valor R$ (recalculado)", "Racional/limite"])
    for r in [
        ["Materialidade global — ativo consolidado", N_ASSETS, 0.0075, "=B2*C2", materiality["overall_brl"], "0,75% do ativo; julgamento didático para incorporadora intensiva em ativos/estoques."],
        ["Materialidade de execução", "=D2", 0.70, "=B3*C3", materiality["performance_brl"], "70% da global, considerando riscos de receita, estoque e estimativas."],
        ["Materialidade específica — estoques", INVENTORY_TOTAL, 0.005, "=B4*C4", materiality["specific_brl"], "0,50% da maior rubrica de ativo; premissa didática para riscos de existência, custo, VLR e classificação."],
        ["Limiar claramente trivial", "=D2", 0.05, "=B5*C5", materiality["clearly_trivial_brl"], "5% da global como premissa didática; natureza qualitativa pode prevalecer."],
    ]: ws.append(r)
    style_sheet(ws)
    for r in range(2, 6): ws.cell(r, 3).number_format = "0.00%"

    ws = wb.create_sheet("Conciliação")
    ws.append(["Linha de conciliação", "Saldo público R$ mil", "Saldo público R$", "Fonte/página", "Soma sintética R$", "Diferença R$", "Conclusão/limitação"])
    recs = [
        ["Estoques consolidados — total", 6_951_508, INVENTORY_TOTAL, "DFP 2025, BP pp.16–17 e Nota 5 p.69", total_cents/100, "=E2-C2", "Rubrica de maior saldo do ativo; população sintética integral concilia exatamente."],
        ["Terrenos a incorporar", 5_678_238, LAND_TOTAL, "DFP 2025, Nota 5 p.69", class_cents["TERRENOS_A_INCORPORAR"]/100, "=E3-C3", "Maior componente dos estoques; representa 81,7% da rubrica."],
        ["Imóveis em construção", 1_174_207, WIP_TOTAL, "DFP 2025, Nota 5 p.69", class_cents["IMOVEIS_EM_CONSTRUCAO"]/100, "=E4-C4", "Classe pública e população sintética conciliam; itens individuais são fictícios."],
        ["Imóveis concluídos", 99_063, COMPLETED_TOTAL, "DFP 2025, Nota 5 p.69", class_cents["IMOVEIS_CONCLUIDOS"]/100, "=E5-C5", "Classe pública e população sintética conciliam; itens individuais são fictícios."],
        ["Estoques circulantes", 2_086_835, INVENTORY_CURRENT, "DFP 2025, BP p.16; terrenos correntes derivados da Nota 5 menos não circulante", current_cents/100, "=E6-C6", "Distribuição sintética baseada nos saldos consolidados publicados."],
        ["Estoques não circulantes", 4_864_673, INVENTORY_NONCURRENT, "DFP 2025, BP p.17; Nota 5 p.69", noncurrent_cents/100, "=E7-C7", "Todo o saldo não circulante do estoque é terrenos a incorporar, conforme Nota 5."],
    ]
    for r in recs: ws.append(r)
    style_sheet(ws)

    ws = wb.create_sheet("Amostragem PT530")
    ws.append(["Métrica", "Resultado", "Base/observação"])
    measures = [
        ("Conta/população", "Estoques", "Maior rubrica de ativo; Nota 5 detalha três classes."),
        ("População", n, "500.000 itens sintéticos únicos, distribuídos por três classes."),
        ("Saldo da população", total_cents/100, "R$; soma exata à rubrica Estoques da DFP 2025, sem representar razão real."),
        ("Amostra aleatória simples sem reposição", len(sample), f"Seed {seed}; demonstração somente com itens sintéticos."),
        ("Parâmetros do cálculo ilustrativo", "95% confiança; p=50%; precisão=±5 p.p.", "n0 = Z²·p·(1−p)/e²; Z=1,96."),
        ("n mínimo com correção finita", n_required_attribute_ceiling, f"n0={n0_attribute:.2f}; n corrigido={n_required_attribute:.2f}; N={n}."),
        ("Fração amostral (%)", len(sample)/n*100, "N = 500.000."),
        ("Amostra dirigida de maiores valores", len(top_rows), "200 maiores valores; não projetável; separada da SRS."),
        ("Projeção didática de shortfall NRV", projected_shortfall, "R$; estimador média amostral × N; população sintética apenas."),
        ("Limite inferior IC95%", max(0, projected_shortfall-margin), "R$; intervalo normal aproximado com correção finita."),
        ("Limite superior IC95%", projected_shortfall+margin, "R$; intervalo normal aproximado com correção finita."),
    ]
    for r in measures: ws.append(r)
    for rule, vals in exception_summary.items():
        ws.append([f"Exceção {rule}: inserida", vals["injected"], "Registro determinístico do gerador."])
        ws.append([f"Exceção {rule}: detectada", vals["detected"], "Teste executado em 100% da base sintética."])
        ws.append([f"Exceção {rule}: não detectada", vals["missed"], "Comparação dos IDs implantados e identificados."])
        ws.append([f"Exceção {rule}: inesperada", vals["unexpected"], "Detecção sem injeção conhecida; apenas população artificial."])
    style_sheet(ws)

    ws = wb.create_sheet("Fontes")
    ws.append(["Fonte", "URL", "Documento", "Páginas/trecho", "Exercício", "Acesso", "Uso/limitação"])
    ws.append(["CVM", "https://www.rad.cvm.gov.br/ENET/frmGerenciaPaginaFRE.aspx?NumeroSequencialDocumento=155189&CodigoTipoInstituicao=1", "DFP V1, protocolo 021350DFP311220250100155189-71, exercício encerrado em 31/12/2025", "BP pp.16–17; metadados DFP", "2025", ACCESS_DATE, "Fonte primária regulatória; arquivo integral conferido no RI."])
    ws.append(["RI Direcional", "https://api.mziq.com/mzfilemanager/v2/d/ada9bc2c-f7d0-4359-9eaf-851b679ab788/184f5357-c3a8-12a9-3a85-cdb268303bcb?origin=2", "DFP 4T25 V1", "BP pp.16–17; Nota 5 p.69", "2025", ACCESS_DATE, "Fonte pública da conta, das classes e dos saldos publicados."])
    ws.append(["B3", "https://sistemaswebb3-listados.b3.com.br/listedCompaniesPage/main/21350/DIRR/overview?language=pt-br", "Ficha da companhia", "sem paginação — identidade e dados financeiros", "2025", ACCESS_DATE, "Identidade/ticker e confirmação cruzada do ativo total; não substitui DFP auditada."])
    style_sheet(ws)

    ws = wb.create_sheet("Dicionário"); ws.append(["Variável", "Tipo", "Descrição", "Domínio/unidade"])
    for row in dictionary: ws.append(row)
    style_sheet(ws)
    ws = wb.create_sheet("Amostra SRS"); sample_headers = FIELDS + ["detected_exception_codes", "sample_method"]
    ws.append(sample_headers)
    for row in sample: ws.append([row.get(k, "") for k in sample_headers])
    style_sheet(ws)
    wb.save(outdir / "papel_amostragem_conciliacao.xlsx")
    print(json.dumps({k: v for k, v in summary.items() if k not in ("population_exception_results", "sources")}, ensure_ascii=False, indent=2))


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--input", default="/home/ubuntu/direcional_audit/data/estoques_sinteticos.csv.gz")
    p.add_argument("--outdir", default="/home/ubuntu/direcional_audit/analysis")
    p.add_argument("--sample-size", type=int, default=1000)
    p.add_argument("--seed", type=int, default=20251232)
    a = p.parse_args(); analyze(Path(a.input), Path(a.outdir), a.sample_size, a.seed)


if __name__ == "__main__": main()
