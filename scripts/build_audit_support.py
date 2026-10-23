#!/usr/bin/env python3
"""Build editable audit planning/analytics workbook for the didactic dossier."""
from __future__ import annotations
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
ANALYSIS = ROOT / "analysis"
DATA = ROOT / "data"

def style(ws):
    ws.freeze_panes = "A2"
    if ws.max_row and ws.max_column:
        ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 34
    for col in ws.columns:
        cells = list(col)
        width = min(54, max(12, max((len(str(c.value or "")) for c in cells[:250]), default=10) + 2))
        ws.column_dimensions[get_column_letter(cells[0].column)].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)

def main():
    results = json.loads((ANALYSIS / "PT530_resultados.json").read_text(encoding="utf-8"))
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    ws = wb.active; ws.title = "Análises"
    ws.append(["Métrica", "2024 R$", "2025 R$", "Variação R$", "Variação %", "Expectativa / limite de triagem", "Leitura e referência"])
    annual = [
        ["Ativo consolidado", 10_067_406_000, 13_219_635_000, "=C2-B2", "=D2/B2", "Triagem: >10% ou >materialidade de execução", "Aumento de 31,3%; DFP BP impresso pp.16–17; risco de expansão/integração."],
        ["Estoques consolidados", 5_532_090_000, 6_951_508_000, "=C3-B3", "=D3/B3", "Triagem: >10% ou >R$69,4 milhões", "Aumento de 25,7%; DFP Nota 5 p.69."],
        ["Terrenos a incorporar", 4_587_823_000, 5_678_238_000, "=C4-B4", "=D4/B4", "Maior componente dos estoques; materialidade específica da rubrica Estoques: R$34,8 milhões", "Aumento de 23,8%; representa 81,7% do estoque e 43,0% do ativo em 2025; Nota 5 p.69."],
        ["Receita de vendas e serviços", 3_348_470_000, 4_343_008_000, "=C5-B5", "=D5/B5", "Triagem: >10% ou >R$69,4 milhões", "Aumento de 29,7%; POC e reconhecimento de receita são PAA; DRE consolidada p.21, auditor p.115–116."],
        ["Resultado antes de tributos", 792_298_000, 1_078_495_000, "=C6-B6", "=D6/B6", "Triagem: >10% ou >R$69,4 milhões", "Aumento de 36,1%; analisar itens não recorrentes e projeções; DRE p.21."],
        ["Empréstimos e financiamentos", 1_549_700_000, 2_758_300_000, "=C7-B7", "=D7/B7", "Mudança de dívida >10%; recomputar covenants e maturidade", "Aumento de 78,0%; 90% de longo prazo, prazo médio 64 meses; relatório da administração p.38."],
        ["Caixa e equivalentes + aplicações", 1_615_200_000, 2_182_800_000, "=C8-B8", "=D8/B8", "Conferir saldos bancários e restrições", "Aumento de 35,1%; caixa da administração, relatório p.38; diferenças de escopo de apresentação."],
        ["Posição de swaps", 2_600_000, 43_000_000, "=C9-B9", "=D9/B9", "Examinar valuation e liquidação", "Aumento; administração relata proteção de taxa; relatório p.38."],
        ["Dívida líquida reportada", -68_200_000, 532_600_000, "=C10-B10", "=IF(B10=0,0,D10/ABS(B10))", "Mudança qualitativa e quantitativa", "Sai de caixa líquido para dívida líquida; aumento atribuído pela administração a dividendos extraordinários de R$804 milhões pagos no fim de 2025; p.32 e p.38."],
    ]
    for row in annual: ws.append(row)
    ws.append(["Nota", "", "", "", "", "", "Percentuais/limites são ferramentas didáticas de triagem; variação não é distorção identificada e precisa ser explicada por documento/evidência."])
    style(ws)
    for r in range(2, 11): ws.cell(r, 5).number_format = "0.0%"

    ws = wb.create_sheet("Pós-fechamento")
    ws.append(["Data/período", "Métrica/evento", "Comparativo/quantia", "Leitura de auditoria", "Tratamento para DFP 2025", "Fonte"])
    rows = [
        ["04/03/2026", "Emissão de CRI", "R$437,5 milhões; 7 e 10 anos", "Alongamento/refinanciamento posterior, mas incluir cronologia e condições de dívida", "Evento não ajustável, divulgado na Nota 26; verificar se revela condições existentes em 31/12/2025", "DFP 2025, Nota 26, impresso p.114; Administração p.38"],
        ["30/06/2026", "Ativo total consolidado", "R$14.628.318 mil; +10,7% desde 31/12/2025", "Crescimento do ativo; avaliar financiamento, capital de giro e integração", "Contexto posterior; não altera o balanço de 31/12/2025 por si só", "ITR 2T26, divulgação de resultados, impresso p.22"],
        ["30/06/2026", "Terrenos a incorporar", "R$6.345.109 mil versus R$5.678.238 mil", "Aumento de R$666.871 mil; risco de mensuração, cronograma e custo de capitalização", "Subsequente; avaliar fatos que evidenciem condição no encerramento sem presumir ajuste", "ITR 2T26, Nota 5, impressos pp.57–58"],
        ["30/06/2026", "Ajuste de estoques a valor realizável", "R$4.310 mil, ante R$4.165 mil; adições R$397 mil e reversões R$252 mil", "Reforça necessidade de testar método, dados, sensibilidade e disponibilidade de informação por projeto", "Evidência posterior; não extrapolar para o saldo de 2025", "ITR 2T26, Nota 5, impressos pp.57–58"],
        ["30/06/2026", "Caixa, equivalentes e aplicações", "R$2.750.097 mil (corrente + aplicações não correntes)", "Indicador favorável de liquidez reportado, mas não substitui maturidade, restrições e caixa projetado", "Contexto posterior; separar de teste de continuidade a 31/12/2025", "ITR 2T26, quadro consolidado, impresso p.22"],
        ["3T26, divulgado 14/10/2026", "Vendas líquidas e lançamentos", "R$1,302 bi; -21,4% vs 2T26 e -20,7% vs 3T25; lançamentos R$1,3525 bi", "Prévia não auditada menciona menor conversão, concentração no fim do trimestre e distratos acima do normal", "Sinal posterior para atualizar riscos, não evidência de condição de 2025 sem ligação temporal", "Prévia operacional 3T26, pp.1–4"],
        ["25/09 a 05/10/2026", "16ª emissão de debêntures", "R$150 milhões; oferta encerrada em 05/10", "Evento de financiamento; conferir finalidade, juros, garantias, covenants e maturidade", "Posterior à emissão da DFP 2025, não ajustável por si só", "Anúncio RI de início e encerramento, 01/10 e 05/10/2026"],
        ["15–16/10/2026", "Participação acionária", "Driehaus 4,7805% do capital; sem intenção declarada de mudar controle", "Evento de acionista, não constitui ativo/passivo da companhia; confirmar apenas se evento de governança material", "Não ajustável à DFP 2025", "Comunicado RI de 16/10/2026"],
        ["09/03 vs 10/03/2026", "Datas do relatório/aprovação", "PwC datado 09/03; administração informa aprovação 10/03", "Inconsistência cronológica aparente; solicitar ata, versão assinada e sequência de aprovação/relatório", "Ponto documental a esclarecer; sem conclusão de irregularidade", "Relatório PwC p.117; Nota 1 p.41"],
    ]
    for r in rows: ws.append(r)
    style(ws)

    ws = wb.create_sheet("Materialidade")
    ws.append(["Nível", "Base R$", "Percentual", "Fórmula R$", "Resultado recalculado R$", "Justificativa/limite"])
    mat = results["materiality"]
    rows = [
        ["Global", 13_219_635_000, 0.0075, "=B2*C2", mat["overall_brl"], "0,75% do ativo consolidado; julgamento didático para incorporadora intensiva em ativos."],
        ["Execução", "=D2", 0.70, "=B3*C3", mat["performance_brl"], "70% da global; risco elevado em estoques, receita e estimativas."],
        ["Específica — Estoques", 6_951_508_000, 0.005, "=B4*C4", mat["specific_brl"], "0,50% da maior rubrica do ativo, conforme Nota 5; julgamento didático."],
        ["Claramente trivial", "=D2", 0.05, "=B5*C5", mat["clearly_trivial_brl"], "5% da global como premissa de apresentação; natureza qualitativa pode sobrepor o limiar."],
    ]
    for r in rows: ws.append(r)
    ws.append(["Aviso", "", "", "", "", "Percentuais não são mandatórios, não são a materialidade efetivamente fixada pelo auditor independente e devem ser recalibrados com resultados finais e dispersão de erros."])
    style(ws)
    for r in range(2,6): ws.cell(r,3).number_format="0.00%"

    ws = wb.create_sheet("Matriz de riscos")
    ws.append(["Conta/processo", "Saldo público e peso", "Afirmações", "Risco e origem", "Significativo?", "Resposta planejada", "Evidência pública realizada / lacuna"])
    risk_rows = [
        ["Terrenos a incorporar", "R$5.678,238 mi; 81,7% do estoque; 43,0% do ativo", "Existência, direitos e obrigações, avaliação, corte, classificação, apresentação", "Título/matrícula, permutas, custo capitalizado, NRV, lançamento planejado e concentração de SPEs", "Sim — montante e complexidade", "Testar contratos e escrituras; confirmar direitos/ônus; inspeção física por amostra; recomputar custo e permutas; reexecutar NRV/sensibilidade; testar classificação e corte", "Tie-out público à Nota 5 executado. Registro auxiliar, título, avaliações e inspeções não acessíveis; população apenas sintética."],
        ["Receita de imóveis em construção (POC)", "Receita R$4.343,008 mi; resultado antes de tributos R$1.078,495 mi", "Ocorrência, exatidão, corte, integridade, classificação", "Reconhecimento depende de contratos, estágio de obra, custo incorrido/orçado e estimativas; PAA PwC e presunção de risco de fraude", "Sim", "Recalcular POC por empreendimento; reconciliar contratos, orçamento, custo, medições e subledger; testar lançamentos manuais/corte; inspeções e confirmação de vendas", "DRE/nota/PAA lidos. Contratos, orçamento, medições e razão indisponíveis; nenhum recálculo interno efetuado."],
        ["Contas a receber e adiantamentos", "Contas a receber consolidadas R$1.438,700 mi; receita imobiliária com POC", "Existência, direitos, avaliação, integridade, corte", "PECLD, distratos, atualização monetária, parcelas recebidas e POC", "Sim", "Circularizar clientes e agentes; testar recebimentos subsequentes, aging, distratos, taxas efetivas e PECLD; recalcular saldo POC", "Saldo e política pública lidos. Confirmações e testes subsequentes não executados por falta de dados."],
        ["Caixa, dívida, swaps e continuidade", "Empréstimos/financiamentos R$2.758,3 mi; dívida líquida R$532,6 mi; alavancagem 23%", "Existência, direitos, integridade, classificação e apresentação", "Pagamento extraordinário de dividendos R$804 mi; fluxo operacional informado R$104 mi; serviço da dívida, swaps, covenants e refinanciamento", "Sim", "Confirmações bancárias; examinar contratos, vencimentos, covenants, swaps, restrições; testar caixa pós-fechamento e modelo de liquidez", "Informação pública lida; não há extrato, covenant certificate ou forecast da administração para testar."],
        ["Provisões para contingências", "PAA PwC; quantificação/rollforward a conferir em Nota 16.2", "Integridade, avaliação, classificação, divulgação", "Prognóstico jurídico, complexidade processual e julgamento da administração", "Sim — PAA", "Cartas a advogados; rollforward por causa; confrontar probabilidade/perda/valor; rever eventos subsequentes e sensibilidade", "PAA e nota pública lidas; nenhuma circularização jurídica ou revisão interna realizada."],
        ["Partes relacionadas", "Ativo R$92,922 mi; passivo R$214,534 mi", "Integridade, ocorrência, exatidão, direitos, divulgação", "SPE/SCP, mútuos sem juros, conta-corrente, garantias, operações fora do curso normal", "Sim — risco qualitativo/completude", "Reconciliar cadastro e declarações de diretores; inspecionar atas/contratos; testar termos, liquidação, garantias e divulgações", "Nota 6 pública lida; cadastro/atas/razão de grupo indisponíveis."],
        ["Grupo e consolidação", "Controladas, SPE/SCP, controladas em conjunto e coligadas", "Completude, direitos, avaliação, consolidação, apresentação", "Perímetro, controle/controle conjunto, eliminações, participações não controladoras e auditores de componentes", "Sim", "Obter organograma e trial balances; testar consolidação/eliminações/políticas; avaliar materialidade de componente e instruções/reportes", "DFP declara uso de controladas/SPE/SCP; não foi possível reconstruir perímetro integral e evidência de componente por fontes públicas."],
        ["Estoques de imóveis em construção e juros capitalizados", "Consolidado R$1.174,207 mi em construção; juros capitalizados na política", "Existência, avaliação, corte, apresentação", "Custos de obra e taxa de capitalização, custo para completar, venda pós-período e estoque de projetos", "Sim", "Confirmar estoque físico; recalcular juros elegíveis/taxa; testar custo incorrido e orçamento; custo versus NRV por empreendimento", "Saldo/política pública; sem inspeção, orçamento ou planilha de juros."],
        ["Imobilizado", "Imobilizado consolidado conforme Balanço p.17", "Existência, direitos, avaliação, classificação", "Adições, baixas, depreciação, impairment e ativos próprios/empresas do grupo", "A avaliar após materialidade por classe", "Reconciliar razão; inspecionar adições e ativos; testar depreciação, titularidade e impairment", "Saldo público; razão, notas fiscais e verificação física não disponíveis."],
    ]
    for r in risk_rows: ws.append(r)
    style(ws)

    ws = wb.create_sheet("Programas de auditoria")
    ws.append(["Ciclo", "Código", "Procedimento / objetivo", "Afirmações", "Status", "Evidência / referência", "Tratamento"])
    programs = [
        ["Receitas", "REV-01", "Reconciliar receita por mês e empreendimento a DRE/razão; comparar venda, POC, custos incorridos/orçados e margem.", "Ocorrência, exatidão, corte, integridade", "Planejado por falta de acesso", "DFP DRE p.21; Nota de receitas; PAA PwC p.115–116", "Obter base de contratos/orçamentos, recalcular POC e explicar desvios; risco de fraude presumido."],
        ["Receitas", "REV-02", "Selecionar contratos e distratos antes/depois de 31/12; recalcular preço, aditivos, saldo e receita cumulativa.", "Ocorrência, exatidão, corte", "Planejado por falta de acesso", "Nota 2.1.7 e relatório PwC p.115–116", "Inspecionar contratos, aditivos e recebimentos; não houve confirmação ou inspeção real."],
        ["Receitas", "REV-03", "Testar lançamentos manuais perto do fechamento e lançamentos fora do padrão; analisar usuários, contrapartidas, horário e estornos.", "Ocorrência, integridade, corte", "Planejado por falta de acesso", "PT240/315 e DRE pública", "Solicitar razão detalhado, logs e parâmetros; procedimento não executado sem sistema."],
        ["Disponibilidades", "CASH-01", "Ler e cotejar o saldo de caixa/equivalentes publicado no balanço com a divulgação de endividamento da administração.", "Existência, apresentação e classificação — apenas contexto público", "Realizado com evidência pública", "Balanço p.16; Administração p.38", "Cotejo documental concluído; não é conciliação a razão/extrato e não comprova existência ou disponibilidade."],
        ["Disponibilidades", "CASH-02", "Reconciliar contas bancárias ao razão e aos extratos; obter confirmações bancárias sob controle do auditor e investigar diferenças.", "Existência, direitos, exatidão, corte", "Planejado por falta de acesso", "Balanço p.16; Administração p.38", "Obter razão, extratos e conciliações; nenhuma confirmação foi enviada."],
        ["Disponibilidades", "CASH-03", "Inspecionar aplicações restritas, títulos e swaps; recalcular valor justo, juros e liquidação subsequente.", "Direitos, avaliação, apresentação", "Planejado por falta de acesso", "Notas financeiras e nota de dívida/swaps", "Obter contratos, custódia, cotações, confirmação independente e pagamentos."],
        ["Estoques", "INV-01", "Conciliar Nota 5 e classes circulante/não circulante; analisar aumento, composição, custo e NRV divulgado.", "Exatidão, classificação, avaliação, apresentação", "Realizado com evidência pública", "DFP Nota 5 p.69; Balanço p.16–17", "Terrenos a incorporar R$5.678.238 mil; conferência pública de soma efetuada."],
        ["Estoques", "INV-02", "Selecionar terrenos e imóveis; examinar contrato, escritura/matrícula, ônus, pagamentos, permuta e cut-off.", "Existência, direitos, avaliação, corte", "Planejado por falta de acesso", "Nota 2.1.8/Nota 5", "Sem livros, contratos, registros e acesso físico. Não alegado como feito."],
        ["Estoques", "INV-03", "Reexecutar NRV por projeto/unidade; comparar custos e custos para completar com preço de venda realizável, distratos e venda subsequente; sensibilidade.", "Avaliação e apresentação", "Planejado por falta de acesso", "Nota 5; PAA de receita", "Obter orçamento, laudos, VGV/vendas e distratos; simulação quantitativa limitada à base sintética PT530."],
        ["Estoques", "INV-04", "Executar regras analíticas, reconciliação e amostra sobre a população sintética de terrenos.", "Exatidão, avaliação, classificação", "Simulado para fins didáticos", "Base CSV.gz; PT530; planilha de conciliação", "500.000 linhas únicas; achados são exclusivamente sintéticos, não achados da companhia."],
        ["Imobilizado", "PPE-01", "Reconciliar cadastro/razão e inspecionar itens selecionados; testar adições, baixas, existência, título, depreciação e impairment.", "Existência, direitos, avaliação, classificação", "Planejado por falta de acesso", "Balanço p.17; Nota imobilizado", "Obter cadastro, notas fiscais, contratos e observação física; nenhum teste feito."],
        ["Imobilizado", "PPE-02", "Inspecionar juros capitalizados a imobilizado/ativos de qualificação e data de disponibilidade para uso.", "Avaliação, corte, classificação", "Planejado por falta de acesso", "Políticas contábeis", "Recalcular taxa/base/tempo e cessar capitalização quando disponível."],
        ["Passivos", "LIAB-01", "Realizar busca de passivos não registrados por pagamentos posteriores, notas fiscais, extratos, contratos, fornecedores e atas.", "Integridade, corte, classificação", "Planejado por falta de acesso", "Balanço p.18–20; notas de dívida/fornecedores", "Examinar pagamentos/subsequent disbursements e confirmar credores; não executado."],
        ["Passivos", "LIAB-02", "Reconciliar dívida/cessões, confirmar credores, testar juros, vencimentos, covenants, garantias, swaps e classificação corrente.", "Integridade, obrigação, exatidão, classificação", "Planejado por falta de acesso", "Administração p.38; notas 9 e 10", "Dívida pública: R$2.758,3 mi; obter instrumentos/confirmar termos antes de concluir."],
        ["Passivos", "LIAB-03", "Confirmar saldo e termos com partes relacionadas; testar mútuos, garantias, liquidações e divulgação.", "Integridade, apresentação, ocorrência", "Planejado por falta de acesso", "Nota 6 p.70", "Saldos consolidados 2025: ativo R$92,922 mi e passivo R$214,534 mi."],
        ["Estimativas", "EST-01", "Recalcular e sensibilizar custos incorridos/custo a completar e receita POC por SPE/empreendimento.", "Avaliação, exatidão, corte", "Planejado por falta de acesso", "PAA PwC p.115–116", "Testar premissas e viés retrospectivo; dados internos ausentes."],
        ["Estimativas", "EST-02", "Reexecutar NRV por terreno/unidade, comparar preços pós-período, distratos e impairment; sensibilizar custo de conclusão e taxa.", "Avaliação e apresentação", "Planejado por falta de acesso", "Nota 5 p.69; ITR 2T26 Nota 5", "Ajustes públicos R$4,165 mi (2025), R$4,310 mi (30/06/26); não é conclusão de adequação."],
        ["Estimativas", "EST-03", "Obter cartas de advogados e testar probabilidade/valor dos processos; reexecutar rollforward das provisões.", "Integridade, avaliação, apresentação", "Planejado por falta de acesso", "PAA PwC p.116–117; Nota 16.2", "Não houve circularização por este trabalho."],
    ]
    for r in programs: ws.append(r)
    style(ws)

    ws = wb.create_sheet("Distorções sintéticas")
    ws.append(["Exceção implantada", "Inserida", "Detectada", "Não detectada", "Inesperada", "Exposição/efeito", "Corrigida?", "Conclusão e distinção"])
    effects = {
        "TITLE_DOCUMENT_INCOMPLETE": "Direitos/evidência insuficientes no registro sintético; não é estimativa de perda.",
        "AREA_INVALID": "A área do registro sintético é zero; efeito monetário não estimado sem método de mensuração.",
        "NRV_BELOW_BOOK": "Shortfall sintético agregado conhecido: R$1.026.809,57; cenário didático, não perda real.",
        "CLASSIFICATION_MISMATCH": "Erro sintético entre classificação e data de lançamento; sem inferência sobre apresentação da companhia.",
        "VALUATION_STALE": "Data de avaliação sintética desatualizada; exposição a estimativa, sem distorção monetária medida.",
    }
    for rule, values in results["population_exception_results"].items():
        ws.append([rule, values["injected"], values["detected"], values["missed"], values["unexpected"], effects[rule], "Não se registra correção em livros reais. As exceções só existem no CSV sintético.", "Testes de regras identificaram todas as exceções implantadas e não marcaram casos não implantados dentro da base artificial."])
    ws.append(["Datas documentais DFP 2025", 1, 1, 0, 0, "Relatório PwC p.117: 09/03/2026; Nota 1 p.41: aprovação em 10/03/2026. Inconsistência cronológica aparente; causa e efeito não quantificados.", "Pendente de esclarecimento; nenhuma correção efetuada", "Não é distorção financeira identificada; solicitar ata/versão assinada e cronologia ao responsável adequado. Não concluir irregularidade nem deficiência sem evidência adicional."])
    style(ws)

    ws = wb.create_sheet("Referências")
    ws.append(["ID", "Fonte", "Documento", "URL", "Páginas/trecho", "Data de acesso", "Uso/observação"])
    refs = [
        ["S1", "CVM", "DFP V1 31/12/2025, NSD 155189", "https://www.rad.cvm.gov.br/ENET/frmGerenciaPaginaFRE.aspx?NumeroSequencialDocumento=155189&CodigoTipoInstituicao=1", "Metadados/protocolo", "22/10/2026", "Fonte regulatória primária."],
        ["S2", "RI Direcional", "DFP 4T25 V1", "https://api.mziq.com/mzfilemanager/v2/d/ada9bc2c-f7d0-4359-9eaf-851b679ab788/184f5357-c3a8-12a9-3a85-cdb268303bcb?origin=2", "BP pp.16–17; Nota 5 p.69; segmentos p.44; partes relacionadas pp.70–71; contingências pp.93–94; Administração pp.32/38; auditor pp.115–117", "22/10/2026", "PDF direto examinado; índice institucional em FONTES.md."],
        ["S3", "B3", "Ficha Direcional Engenharia S.A.", "https://sistemaswebb3-listados.b3.com.br/listedCompaniesPage/main/21350/DIRR/overview?language=pt-br", "Sem página; dados da companhia", "22/10/2026", "Identidade/ticker e ativo total."],
        ["S4", "RI Direcional", "DFP 4T24 V1", "https://api.mziq.com/mzfilemanager/v2/d/ada9bc2c-f7d0-4359-9eaf-851b679ab788/71cc6d14-36b4-e7cc-a1e3-7f7974c42ed7?origin=2", "Relatório auditor pp.111–113", "22/10/2026", "Opinião anterior sem modificação; suporte público a PT510."],
        ["S5", "RI Direcional", "ITR 2T26 V1", "https://api.mziq.com/mzfilemanager/v2/d/ada9bc2c-f7d0-4359-9eaf-851b679ab788/e7250290-a0ce-9144-cc4e-169db46844eb?origin=2", "BP p.22; Nota 5 pp.57–58", "22/10/2026", "Contexto posterior; informações intermediárias."],
        ["S6", "RI Direcional", "Prévia operacional 3T26", "https://filemanager-cdn.mziq.com/published/ada9bc2c-f7d0-4359-9eaf-851b679ab788/ec557d57-e4bc-4648-bb3f-e272f7f0ba0a", "pp.1–4", "22/10/2026", "Prévia não auditada, pós-data."],
        ["S7", "RI Direcional", "Anúncio de início da 16ª emissão de debêntures", "https://filemanager-cdn.mziq.com/published/ada9bc2c-f7d0-4359-9eaf-851b679ab788/2d709687-6743-41d0-a74a-cc287ddf5167", "pp.1–5", "22/10/2026", "Emissão quirografária de R$150 milhões, 25/09/2026."],
        ["S8", "RI Direcional", "Anúncio de encerramento da 16ª emissão de debêntures", "https://filemanager-cdn.mziq.com/published/ada9bc2c-f7d0-4359-9eaf-851b679ab788/c36f23cd-1a25-423a-806b-31f67912116f", "pp.1–5", "22/10/2026", "Oferta encerrada em 05/10/2026."],
        ["S9", "RI Direcional", "Comunicado sobre participação acionária de 16/10/2026", "https://filemanager-cdn.mziq.com/published/ada9bc2c-f7d0-4359-9eaf-851b679ab788/15fa5da0-e732-4b16-b364-0f264c093d07", "p.1", "22/10/2026", "Participação de 4,7805%; sem intenção declarada de mudar controle."],
    ]
    for r in refs: ws.append(r)
    style(ws)

    ws = wb.create_sheet("Notas de uso")
    ws.append(["Tema", "Nota"])
    for r in [
        ["Materialidade", "Percentuais e valores são julgamento didático independente, não a materialidade do auditor PwC."],
        ["Base sintética", "500.000 linhas com identificadores únicos; saldo total fecha com Nota 5. A classificação corrente/não corrente também fecha por desenho com R$813.565 mil/R$4.864.673 mil."],
        ["Exceções", "Erros deliberadamente implantados para testar controles e análises. O registro de injeção fica em data/injection_register.json."],
        ["Procedimentos", "Procedimentos reais de confirmação, inspeção física, entrevista, reperformance interna e acesso a sistema não foram realizados. Estão classificados como planejados por falta de acesso."],
        ["Documentos 2026", "ITR 2T26 e prévia 3T26 são informação posterior; a prévia é expressamente não auditada. Usados apenas para avaliação de eventos subsequentes e continuidade, não para reabrir valores de 2025 sem prova de condição existente na data-base."],
        ["Inconsistência de datas", "Relatório PwC impresso p.117 traz 09/03/2026; Nota 1 p.41 informa aprovação em 10/03/2026. Ponto documental a esclarecer, não conclusão de irregularidade."],
    ]: ws.append(r)
    style(ws)
    out = ANALYSIS / "analises_e_programas.xlsx"
    wb.save(out)
    print(f"Arquivo criado: {out} ({out.stat().st_size:,} bytes)")

if __name__ == "__main__": main()
